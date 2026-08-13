#!/usr/bin/env python3
"""Convert PROJECT_FILES_GUIDE.md to DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parent.parent
MD_PATH = PROJECT_ROOT / "docs" / "PROJECT_FILES_GUIDE.md"
DOCX_PATH = PROJECT_ROOT / "docs" / "PROJECT_FILES_GUIDE.docx"


def set_doc_styles(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for level in range(1, 4):
        heading = doc.styles[f"Heading {level}"]
        heading.font.name = "Calibri"
        heading.font.color.rgb = RGBColor(0x1F, 0x47, 0x88)


def add_formatted_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def parse_table_rows(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if not line.strip().startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(set(c) <= {"-", ":", " "} for c in cells):
            continue
        rows.append(cells)
    return rows


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            add_formatted_runs(p, cell_text)
            if r_idx == 0:
                for run in p.runs:
                    run.bold = True


def convert_md_to_docx(md_path: Path, docx_path: Path) -> None:
    content = md_path.read_text(encoding="utf-8")
    lines = content.splitlines()
    doc = Document()
    set_doc_styles(doc)

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    i = 0
    in_code = False
    code_lines: list[str] = []
    code_lang = ""

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lang = line.strip()[3:].strip()
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.style = doc.styles["Normal"]
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                pf = p.paragraph_format
                pf.left_indent = Inches(0.25)
                pf.space_before = Pt(6)
                pf.space_after = Pt(6)
                in_code = False
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.strip() == "---":
            doc.add_paragraph()
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
            i += 1
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue

        if line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, parse_table_rows(table_lines))
            doc.add_paragraph()
            continue

        if line.strip().startswith("- [ ]"):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p, line.strip()[5:].strip())
            i += 1
            continue

        if line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p, line.strip()[2:].strip())
            i += 1
            continue

        if re.match(r"^\d+\.\s", line.strip()):
            p = doc.add_paragraph(style="List Number")
            add_formatted_runs(p, re.sub(r"^\d+\.\s", "", line.strip()))
            i += 1
            continue

        if line.strip().startswith("> "):
            p = doc.add_paragraph()
            add_formatted_runs(p, line.strip()[2:].strip())
            p.paragraph_format.left_indent = Inches(0.35)
            i += 1
            continue

        if line.strip().startswith("*") and line.strip().endswith("*"):
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = p.add_run(line.strip().strip("*"))
            run.italic = True
            i += 1
            continue

        if line.strip():
            p = doc.add_paragraph()
            add_formatted_runs(p, line.strip())
        i += 1

    doc.save(docx_path)
    print(f"Created: {docx_path}")


if __name__ == "__main__":
    convert_md_to_docx(MD_PATH, DOCX_PATH)
